"""
OMR Scanner — Flask App
"""

import os
import uuid
import base64
import json
import cv2
import numpy as np
from io import BytesIO
from flask import Flask, request, jsonify, render_template, send_file, Response
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from omr_engine import process_omr, generate_sample_sheet, OMRConfig

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
ALLOWED    = {"png", "jpg", "jpeg", "webp", "bmp"}

app = Flask(__name__)
app.config["UPLOAD_FOLDER"]      = UPLOAD_DIR
app.config["MAX_CONTENT_LENGTH"] = 300 * 1024 * 1024
os.makedirs(UPLOAD_DIR, exist_ok=True)


def allowed(f):
    return "." in f and f.rsplit(".", 1)[1].lower() in ALLOWED

def img_b64(img):
    ok, buf = cv2.imencode(".png", img)
    return base64.b64encode(buf.tobytes()).decode() if ok else ""

def parse_key(raw):
    key = json.loads(raw)
    if len(key) != OMRConfig.NUM_QUESTIONS:
        raise ValueError(f"Answer key needs {OMRConfig.NUM_QUESTIONS} entries")
    if any(k not in (0,1,2,3) for k in key):
        raise ValueError("Each answer must be 0-3 (A-D)")
    return key


# ── Excel builder ─────────────────────────────────────────────────────────────

def build_excel(results, answer_key):
    wb  = Workbook()
    lbl = OMRConfig.CHOICE_LABELS

    hdr_fill  = PatternFill("solid", fgColor="1F3864")
    ok_fill   = PatternFill("solid", fgColor="C6EFCE")
    bad_fill  = PatternFill("solid", fgColor="FFC7CE")
    amb_fill  = PatternFill("solid", fgColor="FFEB9C")
    alt_fill  = PatternFill("solid", fgColor="EEF2FF")
    hdr_font  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    body_font = Font(name="Arial", size=10)
    bold_font = Font(bold=True, name="Arial", size=10)
    thin      = Side(style="thin", color="CCCCCC")
    bdr       = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr       = Alignment(horizontal="center", vertical="center")

    # ── Summary sheet ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    heads = ["#","Filename","Chest No","Company","Score","Total","%","Correct","Wrong","Blank","Grade"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = ctr; cell.border = bdr
    ws.row_dimensions[1].height = 22

    for i, r in enumerate(results, 2):
        pct   = r["score_pct"]
        grade = ("A+" if pct>=90 else "A" if pct>=80 else "B" if pct>=70
                 else "C" if pct>=60 else "D" if pct>=50 else "F")
        rf    = alt_fill if i%2==0 else PatternFill("solid", fgColor="FFFFFF")
        vals  = [i-1, os.path.basename(r["filename"]), r["student_id"],
                 r.get("company","?"),
                 r["score"], r["total"], pct,
                 r["correct"], r["wrong"], r["blank"], grade]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font=body_font; cell.alignment=ctr; cell.border=bdr; cell.fill=rf
        pct_cell = ws.cell(row=i, column=7)
        pct_cell.number_format = '0.0"%"'
        pct_cell.fill = ok_fill if pct>=60 else (amb_fill if pct>=40 else bad_fill)

    n = len(results)
    if n:
        sr = n + 3
        for k, (label, fn) in enumerate([("Average %", "AVERAGE"),
                                         ("Highest %", "MAX"),
                                         ("Lowest %",  "MIN")]):
            ws.cell(row=sr+k, column=2, value=label).font = bold_font
            c = ws.cell(row=sr+k, column=3, value=f"={fn}(G2:G{n+1})")
            c.font = body_font; c.number_format = '0.0"%"'; c.alignment = ctr

    for c, w in enumerate([5,30,12,10,8,8,10,9,8,8,8], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # ── Per-question detail ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Per-Question Detail")
    for col, val in [(1,"Q#"),(2,"Answer Key")]:
        cell = ws2.cell(row=1, column=col, value=val)
        cell.font=hdr_font; cell.fill=hdr_fill; cell.alignment=ctr
    for s, r in enumerate(results, 3):
        cell = ws2.cell(row=1, column=s, value=f"Chest:{r['student_id']}")
        cell.font=hdr_font; cell.fill=hdr_fill; cell.alignment=ctr

    for q in range(OMRConfig.NUM_QUESTIONS):
        row = q+2
        ws2.cell(row=row, column=1, value=q+1).alignment = ctr
        key_lbl = lbl[answer_key[q]] if q < len(answer_key) else "?"
        kc = ws2.cell(row=row, column=2, value=key_lbl)
        kc.font=bold_font; kc.alignment=ctr
        for s, r in enumerate(results, 3):
            bd   = r["breakdown"][q]
            cell = ws2.cell(row=row, column=s, value=bd["detected"])
            cell.font=body_font; cell.alignment=ctr
            cell.fill = (ok_fill if bd["status"]=="correct"
                         else bad_fill if bd["status"]=="wrong" else amb_fill)

    ws2.column_dimensions["A"].width = 5
    ws2.column_dimensions["B"].width = 12
    for s in range(len(results)):
        ws2.column_dimensions[get_column_letter(s+3)].width = 12

    # ── Answer key sheet ─────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Answer Key")
    for c, v in [(1,"Question"),(2,"Correct Answer")]:
        cell = ws3.cell(row=1, column=c, value=v)
        cell.font=hdr_font; cell.fill=hdr_fill; cell.alignment=ctr
    for q, k in enumerate(answer_key):
        ws3.cell(row=q+2, column=1, value=q+1).alignment = ctr
        cell = ws3.cell(row=q+2, column=2, value=lbl[k])
        cell.alignment=ctr; cell.fill=ok_fill
    ws3.column_dimensions["A"].width=12
    ws3.column_dimensions["B"].width=16

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# ── Routes ────────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html", num_questions=OMRConfig.NUM_QUESTIONS)

@app.route("/scan_single", methods=["POST"])
def scan_single():
    if "sheet" not in request.files:
        return jsonify({"error":"No file uploaded"}), 400
    file = request.files["sheet"]
    if not file.filename or not allowed(file.filename):
        return jsonify({"error":"Invalid file"}), 400
    try:
        key = parse_key(request.form.get("answer_key",""))
    except Exception as e:
        return jsonify({"error":str(e)}), 400

    path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}_{secure_filename(file.filename)}")
    file.save(path)
    result = process_omr(path, key)
    result["filename"] = secure_filename(file.filename)
    try: os.remove(path)
    except OSError: pass

    if "error" in result:
        return jsonify(result), 422
    result["annotated_b64"] = img_b64(result.pop("annotated_img"))
    return jsonify(result)

@app.route("/scan_bulk", methods=["POST"])
def scan_bulk():
    files = request.files.getlist("sheets")
    if not files:
        return jsonify({"error":"No files"}), 400
    try:
        key = parse_key(request.form.get("answer_key",""))
    except Exception as e:
        return jsonify({"error":str(e)}), 400

    results = []
    for file in files:
        if not file.filename or not allowed(file.filename):
            continue
        path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}_{secure_filename(file.filename)}")
        file.save(path)
        r = process_omr(path, key)
        r["filename"] = secure_filename(file.filename)
        try: os.remove(path)
        except OSError: pass
        if "error" not in r:
            r["annotated_b64"] = img_b64(r.pop("annotated_img"))
        results.append(r)
    return jsonify({"results": results, "count": len(results)})

@app.route("/export_excel", methods=["POST"])
def export_excel():
    data = request.get_json()
    if not data:
        return jsonify({"error":"No data"}), 400
    buf = build_excel(data["results"], data["answer_key"])
    return send_file(buf, as_attachment=True,
                     download_name="omr_results.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/sample_sheet")
def sample_sheet():
    img = generate_sample_sheet()
    ok, buf = cv2.imencode(".png", img)
    return Response(buf.tobytes(), mimetype="image/png",
                    headers={"Content-Disposition":"attachment;filename=sample_sheet.png"})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", host="0.0.0.0", port=port)
